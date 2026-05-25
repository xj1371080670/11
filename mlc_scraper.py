#!/usr/bin/env python3
"""
MercadoLibre 卖家线索爬虫 - 深圳爱骑仕跨境部
按关键词抓取 MercadoLibre 商品 -> 反查卖家 -> 识别中国卖家 -> 导出 Excel
"""
from __future__ import annotations

import argparse
import json
import os
import random
import re
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Optional
from urllib.parse import quote, urlparse

import requests
import yaml
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from tqdm import tqdm

# ----------------------------------------------------------------------
# 站点配置
# ----------------------------------------------------------------------
SITE_DOMAINS = {
    "MLC": "listado.mercadolibre.cl",   # 智利
    "MLM": "listado.mercadolibre.com.mx",  # 墨西哥
    "MLB": "lista.mercadolivre.com.br",  # 巴西 (注意: mercadolivre 不是 mercadolibre)
    "MLA": "listado.mercadolibre.com.ar",  # 阿根廷
    "MCO": "listado.mercadolibre.com.co",  # 哥伦比亚
}
SITE_HOST = {
    "MLC": "www.mercadolibre.cl",
    "MLM": "www.mercadolibre.com.mx",
    "MLB": "www.mercadolivre.com.br",
    "MLA": "www.mercadolibre.com.ar",
    "MCO": "www.mercadolibre.com.co",
}
SITE_LANG = {
    "MLC": "es-CL,es;q=0.9",
    "MLM": "es-MX,es;q=0.9",
    "MLB": "pt-BR,pt;q=0.9",
    "MLA": "es-AR,es;q=0.9",
    "MCO": "es-CO,es;q=0.9",
}

USER_AGENTS = [
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/121.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/121.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.2 Safari/605.1.15",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:122.0) Gecko/20100101 Firefox/122.0",
]

# ----------------------------------------------------------------------
# 中国卖家识别启发式
# ----------------------------------------------------------------------
CN_LOCATION_HINTS = [
    "china", "shenzhen", "guangzhou", "shanghai", "beijing", "yiwu",
    "ningbo", "hangzhou", "dongguan", "foshan", "xiamen", "fujian",
    "zhejiang", "guangdong", "hong kong", "hk", "taiwan",
    "深圳", "广州", "上海", "北京", "义乌", "宁波", "杭州", "中国", "香港",
]

# 常见汉语拼音姓氏 / 词根 - 用于昵称启发式判断
PINYIN_TOKENS = {
    "zhang", "wang", "li", "chen", "liu", "yang", "huang", "zhao", "wu",
    "zhou", "xu", "sun", "ma", "zhu", "hu", "guo", "he", "gao", "lin",
    "luo", "deng", "feng", "han", "tang", "cao", "peng", "tian", "dong",
    "yu", "jiang", "qin", "shi", "hou", "shen", "xie", "song", "tao",
    "lu", "fu", "lv", "lue", "cui", "duan", "fan", "fang", "ge", "geng",
    # 常见品牌词根
    "shen", "zhen", "guang", "zhou", "yi", "wu", "ke", "ji", "tech",
    "dian", "zi", "tong", "xun", "ai", "qi", "shi",
}

PINYIN_PATTERN = re.compile(
    r"\b(" + "|".join(sorted(PINYIN_TOKENS, key=len, reverse=True)) + r")\b",
    re.IGNORECASE,
)
DIGIT_TAIL = re.compile(r"\d{3,}$")  # 昵称末尾跟 3+ 位数字 (中国卖家常见)


# ----------------------------------------------------------------------
# 数据模型
# ----------------------------------------------------------------------
@dataclass
class Listing:
    keyword: str = ""
    item_id: str = ""
    title: str = ""
    price: Optional[float] = None
    currency: str = ""
    url: str = ""
    sold: Optional[int] = None
    location: str = ""
    seller_id: str = ""
    seller_nickname: str = ""


@dataclass
class Seller:
    seller_id: str = ""
    nickname: str = ""
    country: str = ""
    state: str = ""
    city: str = ""
    address_full: str = ""
    registration_year: Optional[int] = None
    seller_level: str = ""
    sold_total: Optional[int] = None
    is_official_store: bool = False
    profile_url: str = ""
    store_url: str = ""
    likely_chinese: bool = False
    cn_signals: list[str] = field(default_factory=list)
    sample_titles: list[str] = field(default_factory=list)
    listing_count: int = 0


# ----------------------------------------------------------------------
# 工具
# ----------------------------------------------------------------------
def slugify(text: str) -> str:
    text = text.strip().lower()
    # MercadoLibre 把空格 -> '-' , 去除非字母数字
    text = re.sub(r"[\s/]+", "-", text)
    text = re.sub(r"[^\w\-]", "", text, flags=re.UNICODE)
    text = re.sub(r"-+", "-", text).strip("-")
    return text


def make_session(timeout: int) -> requests.Session:
    s = requests.Session()
    s.headers.update({
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Encoding": "gzip, deflate, br",
        "Cache-Control": "no-cache",
        "DNT": "1",
        "Upgrade-Insecure-Requests": "1",
    })
    # 让用户通过环境变量传 HTTPS_PROXY/HTTP_PROXY, requests 自动识别
    s.timeout = timeout
    return s


# 国内常见 VPN 客户端的本地代理端口
COMMON_PROXY_PORTS = [
    ("Clash for Windows / ClashX", 7890),
    ("Clash 备用", 7891),
    ("v2rayN", 10809),
    ("v2rayN SOCKS", 10808),
    ("Shadowsocks", 1087),
    ("Shadowsocks SOCKS", 1080),
    ("Surge", 6152),
    ("通用 8080", 8080),
    ("通用 8888", 8888),
]


def detect_proxy() -> Optional[str]:
    """探测本机正在运行的代理端口, 用一个能确定走代理的请求验证"""
    print("[PROXY] 自动探测本机代理端口...")
    test_url = "https://api.ipify.org?format=json"
    for name, port in COMMON_PROXY_PORTS:
        proxy = f"http://127.0.0.1:{port}"
        try:
            r = requests.get(
                test_url,
                proxies={"http": proxy, "https": proxy},
                timeout=4,
            )
            if r.status_code == 200:
                ip = r.json().get("ip", "?")
                print(f"  ✅ 端口 {port} ({name}) 可用, 出口IP: {ip}")
                return proxy
        except Exception:
            pass
    print("  ❌ 没找到能用的代理端口")
    return None


def verify_proxy_works_for_ml(proxy: str, site: str, timeout: int = 15) -> bool:
    """额外确认代理出口IP能访问 MercadoLibre (没被地理墙拦)"""
    domain = SITE_DOMAINS[site]
    test_url = f"https://{domain}/celular"  # 一个一定有结果的搜索词
    try:
        r = requests.get(
            test_url,
            headers={
                "User-Agent": random.choice(USER_AGENTS),
                "Accept-Language": SITE_LANG[site],
            },
            proxies={"http": proxy, "https": proxy},
            timeout=timeout,
            allow_redirects=True,
        )
        if "/gz/account-verification" in r.url:
            print(f"  ❌ 代理 IP 仍被 MercadoLibre 拦截, 切个海外节点重试")
            return False
        if r.status_code == 200 and "mercadolibre" in r.url:
            print(f"  ✅ 代理可访问 MercadoLibre {site}")
            return True
        print(f"  ⚠️  返回 {r.status_code} {r.url}")
        return False
    except Exception as e:
        print(f"  ❌ {e}")
        return False


def polite_get(
    session: requests.Session,
    url: str,
    *,
    referer: Optional[str] = None,
    lang: str = "es-CL,es;q=0.9",
    timeout: int = 30,
    retries: int = 2,
) -> Optional[str]:
    headers = {
        "User-Agent": random.choice(USER_AGENTS),
        "Accept-Language": lang,
    }
    if referer:
        headers["Referer"] = referer
    for attempt in range(retries + 1):
        try:
            resp = session.get(url, headers=headers, timeout=timeout, allow_redirects=True)
            final_url = resp.url
            # 命中 ML 的反爬验证页
            if "/gz/account-verification" in final_url or "/cgi-bin/captcha" in final_url:
                print(f"  [BLOCKED] 被反爬拦截 -> {final_url}")
                print(f"  请确认你已通过 VPN/代理使用海外 IP 再重试")
                return None
            if resp.status_code == 200:
                return resp.text
            if resp.status_code in (403, 429):
                wait = 5 * (attempt + 1)
                print(f"  [{resp.status_code}] 被限速, 等 {wait}s 后重试...")
                time.sleep(wait)
                continue
            print(f"  [HTTP {resp.status_code}] {url}")
            return None
        except requests.RequestException as e:
            print(f"  [ERR] {e}, 重试 {attempt + 1}/{retries}")
            time.sleep(3 * (attempt + 1))
    return None


def random_sleep(lo: float, hi: float) -> None:
    time.sleep(random.uniform(lo, hi))


# ----------------------------------------------------------------------
# 解析器: 从 HTML 抽取 ML 嵌入的 __PRELOADED_STATE__ JSON
# ----------------------------------------------------------------------
def extract_preloaded_state(html: str) -> Optional[dict]:
    """ML 把页面所有数据塞进一个 inline JSON。优先抓它最可靠。"""
    # Pattern A: <script id="__PRELOADED_STATE__">{...}</script>
    m = re.search(
        r'<script[^>]*id=["\']__PRELOADED_STATE__["\'][^>]*>(.+?)</script>',
        html, re.S
    )
    if m:
        try:
            return json.loads(m.group(1))
        except json.JSONDecodeError:
            pass
    # Pattern B: window.__PRELOADED_STATE__ = {...};
    m = re.search(
        r'window\.__PRELOADED_STATE__\s*=\s*(\{.+?\});\s*</script>',
        html, re.S
    )
    if m:
        try:
            return json.loads(m.group(1))
        except json.JSONDecodeError:
            pass
    return None


def walk_find_all(node: Any, key: str, depth: int = 0) -> Iterable[Any]:
    """递归 yield 出 dict 中所有 key 命中的 value (不深入命中节点本身, 防爆栈)"""
    if depth > 30:
        return
    if isinstance(node, dict):
        for k, v in node.items():
            if k == key:
                yield v
            else:
                yield from walk_find_all(v, key, depth + 1)
    elif isinstance(node, list):
        for v in node:
            yield from walk_find_all(v, key, depth + 1)


def extract_listings_from_state(state: dict, keyword: str) -> list[Listing]:
    """从 __PRELOADED_STATE__ 中找 results 数组并提取每条商品"""
    results: list[dict] = []
    # 找名为 "results" 的 list
    for cand in walk_find_all(state, "results"):
        if isinstance(cand, list) and cand and isinstance(cand[0], dict):
            # 必须看起来像商品数组 - 有 id 或 permalink 字段
            sample = cand[0]
            if any(k in sample for k in ("permalink", "id", "title")):
                results.extend(cand)
    listings: list[Listing] = []
    seen_ids: set[str] = set()
    for r in results:
        if not isinstance(r, dict):
            continue
        item_id = str(r.get("id") or r.get("itemId") or "")
        if not item_id or item_id in seen_ids:
            continue
        seen_ids.add(item_id)

        title = r.get("title") or r.get("name") or ""
        permalink = r.get("permalink") or r.get("url") or ""

        # 价格
        price_obj = r.get("price") or {}
        if isinstance(price_obj, dict):
            price = price_obj.get("amount") or price_obj.get("value")
            currency = price_obj.get("currency_id") or price_obj.get("currency") or ""
        else:
            price = price_obj if isinstance(price_obj, (int, float)) else None
            currency = r.get("currency_id") or ""

        # 销量
        sold = None
        sold_obj = r.get("sold_quantity") or r.get("soldQuantity")
        if isinstance(sold_obj, (int, float)):
            sold = int(sold_obj)
        else:
            # ML 有时返回字符串 "100 vendidos"
            for txt_field in ("subtitle", "soldText"):
                v = r.get(txt_field)
                if isinstance(v, str):
                    m = re.search(r"(\d+[\d,.]*)", v)
                    if m:
                        try:
                            sold = int(m.group(1).replace(",", "").replace(".", ""))
                            break
                        except ValueError:
                            pass

        # 地址
        addr = r.get("seller_address") or r.get("sellerAddress") or {}
        if isinstance(addr, dict):
            city = (addr.get("city") or {}).get("name", "") if isinstance(addr.get("city"), dict) else ""
            state_n = (addr.get("state") or {}).get("name", "") if isinstance(addr.get("state"), dict) else ""
            country = (addr.get("country") or {}).get("name", "") if isinstance(addr.get("country"), dict) else ""
            location = ", ".join(p for p in (city, state_n, country) if p)
        else:
            location = ""

        # 卖家
        seller_obj = r.get("seller") or {}
        seller_id = ""
        seller_nick = ""
        if isinstance(seller_obj, dict):
            seller_id = str(seller_obj.get("id") or seller_obj.get("seller_id") or "")
            seller_nick = seller_obj.get("nickname") or seller_obj.get("eshop", {}).get("nick_name", "") or ""

        listings.append(Listing(
            keyword=keyword,
            item_id=item_id,
            title=str(title),
            price=float(price) if isinstance(price, (int, float)) else None,
            currency=str(currency),
            url=str(permalink),
            sold=sold,
            location=location,
            seller_id=seller_id,
            seller_nickname=str(seller_nick),
        ))
    return listings


def extract_listings_from_html_fallback(html: str, keyword: str) -> list[Listing]:
    """JSON 抓不到时的兜底: 用 BeautifulSoup 解析卡片"""
    soup = BeautifulSoup(html, "lxml")
    listings: list[Listing] = []
    for card in soup.select("li.ui-search-layout__item, div.ui-search-result"):
        a = card.select_one("a.ui-search-link, a.ui-search-result__content")
        if not a:
            continue
        url = a.get("href", "").strip()
        title_el = card.select_one(".ui-search-item__title, h2")
        title = title_el.get_text(" ", strip=True) if title_el else ""
        price_el = card.select_one(".andes-money-amount__fraction")
        price = None
        if price_el:
            try:
                price = float(price_el.get_text(strip=True).replace(".", "").replace(",", "."))
            except ValueError:
                pass
        m = re.search(r"/(MLC|MLM|MLB|MLA|MCO)-?(\d+)", url)
        item_id = (m.group(1) + m.group(2)) if m else ""
        listings.append(Listing(
            keyword=keyword, item_id=item_id, title=title,
            price=price, url=url,
        ))
    return listings


# ----------------------------------------------------------------------
# 卖家详情抓取
# ----------------------------------------------------------------------
def parse_seller_profile(html: str, seller_id: str, nickname: str) -> Seller:
    """解析卖家主页 / 店铺页, 抽出可用信息"""
    s = Seller(seller_id=seller_id, nickname=nickname)

    state = extract_preloaded_state(html)
    if state:
        # 试图找 seller / user 节点
        for cand in walk_find_all(state, "user"):
            if isinstance(cand, dict) and cand.get("id"):
                if str(cand.get("id")) == seller_id or cand.get("nickname") == nickname:
                    s.country = str((cand.get("country_id") or "")) or s.country
                    addr = cand.get("address") or {}
                    if isinstance(addr, dict):
                        s.city = addr.get("city") or s.city
                        s.state = addr.get("state") or s.state
                    reg = cand.get("registration_date")
                    if isinstance(reg, str) and len(reg) >= 4:
                        try:
                            s.registration_year = int(reg[:4])
                        except ValueError:
                            pass
                    break
        for cand in walk_find_all(state, "seller_reputation"):
            if isinstance(cand, dict):
                s.seller_level = cand.get("level_id") or s.seller_level
                tx = cand.get("transactions") or {}
                if isinstance(tx, dict) and isinstance(tx.get("total"), int):
                    s.sold_total = tx["total"]

    # BeautifulSoup 兜底
    soup = BeautifulSoup(html, "lxml")
    page_text = soup.get_text(" ", strip=True).lower()

    # 地址兜底: 找 "Ubicación" / "Localização" 行
    if not s.address_full:
        for label in ("Ubicación", "Localização", "Location"):
            el = soup.find(string=re.compile(label, re.I))
            if el and el.parent:
                # 取兄弟文字
                nxt = el.parent.find_next(string=True)
                if nxt:
                    s.address_full = nxt.strip()[:200]
                    break

    # 累计销量兜底: "X ventas" / "X vendas" / "X sales"
    if s.sold_total is None:
        m = re.search(r"([\d.,]+)\s*(ventas|vendas|sales)\b", page_text)
        if m:
            try:
                s.sold_total = int(m.group(1).replace(".", "").replace(",", ""))
            except ValueError:
                pass

    # 信用等级兜底: MercadoLider Platinum/Gold/...
    if not s.seller_level:
        m = re.search(r"mercadol[ií]der\s*(platinum|gold|silver|plata)", page_text, re.I)
        if m:
            s.seller_level = "MercadoLider " + m.group(1).title()

    # 是否官方店
    if soup.select_one('[class*="official-store"]') or "tienda oficial" in page_text:
        s.is_official_store = True

    return s


# ----------------------------------------------------------------------
# 中国卖家识别
# ----------------------------------------------------------------------
def detect_chinese(seller: Seller, listings_for_seller: list[Listing]) -> tuple[bool, list[str]]:
    signals: list[str] = []
    nick = (seller.nickname or "").lower()
    addr = " ".join(filter(None, [
        seller.country, seller.state, seller.city, seller.address_full,
    ])).lower()
    locs = " ".join(l.location for l in listings_for_seller).lower()
    haystack = f"{addr} {locs}"

    # 信号 1: 地址含中国地名
    for kw in CN_LOCATION_HINTS:
        if kw in haystack:
            signals.append(f"address_contains:{kw}")
            break

    # 信号 2: 昵称含拼音 token
    if PINYIN_PATTERN.search(nick):
        m = PINYIN_PATTERN.search(nick)
        signals.append(f"pinyin_in_nickname:{m.group(1).lower()}")

    # 信号 3: 昵称末尾大量数字 (典型 alibaba/seller 注册号)
    if DIGIT_TAIL.search(nick):
        signals.append("digit_tail_nickname")

    # 信号 4: 昵称包含中国卖家高频词
    for kw in ("china", "shenzhen", "guangzhou", "shop", "store",
               "tech", "trade", "trading", "imp", "exp", "intl"):
        if kw in nick:
            signals.append(f"nickname_kw:{kw}")
            break

    # 综合判定: 命中 1 个强信号 (address) 或 2+ 个弱信号
    has_strong = any(s.startswith("address_contains") for s in signals)
    is_chinese = has_strong or len(signals) >= 2
    return is_chinese, signals


# ----------------------------------------------------------------------
# 主流程
# ----------------------------------------------------------------------
def build_search_url(site: str, keyword: str, page: int) -> str:
    domain = SITE_DOMAINS[site]
    slug = slugify(keyword)
    if page <= 1:
        return f"https://{domain}/{slug}"
    offset = 48 * (page - 1) + 1
    return f"https://{domain}/{slug}_Desde_{offset}_NoIndex_True"


def build_seller_profile_url(site: str, nickname: str) -> str:
    host = SITE_HOST[site]
    return f"https://{host}/perfil/{quote(nickname)}"


def crawl(cfg: dict, args: argparse.Namespace) -> tuple[list[Listing], dict[str, Seller]]:
    site: str = cfg["site"]
    if site not in SITE_DOMAINS:
        raise ValueError(f"Unknown site {site}")

    output_dir = Path(cfg.get("output_dir", "output"))
    output_dir.mkdir(parents=True, exist_ok=True)
    raw_dir = output_dir / "raw"
    if cfg.get("save_raw_html"):
        raw_dir.mkdir(exist_ok=True)

    keywords = [args.keyword] if args.keyword else cfg["keywords"]
    max_pages = args.max_pages or cfg.get("max_pages_per_keyword", 5)
    delay_lo = float(cfg.get("delay_min", 2))
    delay_hi = float(cfg.get("delay_max", 4))
    timeout = int(cfg.get("timeout", 30))
    lang = SITE_LANG[site]

    session = make_session(timeout)

    all_listings: list[Listing] = []
    print(f"\n[CRAWL] site={site} keywords={len(keywords)} max_pages={max_pages}")

    for kw in keywords:
        print(f"\n=== 关键词: {kw} ===")
        for page in range(1, max_pages + 1):
            url = build_search_url(site, kw, page)
            print(f"  P{page} GET {url}")
            html = polite_get(session, url, lang=lang, timeout=timeout)
            if not html:
                print(f"  跳过 P{page}")
                break
            if cfg.get("save_raw_html"):
                fname = raw_dir / f"search_{slugify(kw)}_p{page}.html"
                fname.write_text(html, encoding="utf-8", errors="ignore")

            state = extract_preloaded_state(html)
            if state:
                page_listings = extract_listings_from_state(state, kw)
            else:
                page_listings = extract_listings_from_html_fallback(html, kw)

            if not page_listings:
                print(f"  P{page} 没解析到商品, 终止该关键词")
                break
            print(f"  P{page} 抓到 {len(page_listings)} 条商品")
            all_listings.extend(page_listings)
            random_sleep(delay_lo, delay_hi)

    # 去重 (item_id)
    dedup: dict[str, Listing] = {}
    for l in all_listings:
        if l.item_id and l.item_id not in dedup:
            dedup[l.item_id] = l
    listings = list(dedup.values())
    print(f"\n[LISTINGS] 总计 {len(all_listings)} -> 去重后 {len(listings)}")

    # 按 seller 分组
    by_seller: dict[str, list[Listing]] = {}
    for l in listings:
        if not l.seller_id and not l.seller_nickname:
            continue
        key = l.seller_id or l.seller_nickname
        by_seller.setdefault(key, []).append(l)

    sellers: dict[str, Seller] = {}
    print(f"[SELLERS] 准备处理 {len(by_seller)} 个唯一卖家")

    for key, ls in tqdm(by_seller.items(), desc="抓卖家详情"):
        sample = ls[0]
        seller = Seller(
            seller_id=sample.seller_id,
            nickname=sample.seller_nickname,
            address_full=sample.location,
        )
        seller.profile_url = build_seller_profile_url(site, sample.seller_nickname) if sample.seller_nickname else ""
        seller.listing_count = len(ls)
        seller.sample_titles = [l.title for l in ls[:5]]

        if cfg.get("fetch_seller_detail") and seller.profile_url:
            html = polite_get(session, seller.profile_url, lang=lang, timeout=timeout)
            if html:
                detail = parse_seller_profile(html, seller.seller_id, seller.nickname)
                # 合并: 详情有数据就用详情的
                for fld in ("country", "state", "city", "address_full",
                            "registration_year", "seller_level", "sold_total",
                            "is_official_store"):
                    v = getattr(detail, fld)
                    if v:
                        setattr(seller, fld, v)
                if cfg.get("save_raw_html"):
                    fname = raw_dir / f"seller_{slugify(seller.nickname)}.html"
                    fname.write_text(html, encoding="utf-8", errors="ignore")
            random_sleep(delay_lo, delay_hi)

        is_cn, signals = detect_chinese(seller, ls)
        seller.likely_chinese = is_cn
        seller.cn_signals = signals
        sellers[key] = seller

    return listings, sellers


# ----------------------------------------------------------------------
# Excel 导出
# ----------------------------------------------------------------------
HEADER_FILL = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HOT_FILL = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
WARM_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def autosize(ws, max_width=60):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 8
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(v), max_width))
        ws.column_dimensions[col_letter].width = max_len + 2


def priority_for(seller: Seller) -> str:
    score = 0
    if seller.likely_chinese:
        score += 3
    if (seller.sold_total or 0) >= 1000:
        score += 2
    elif (seller.sold_total or 0) >= 100:
        score += 1
    if seller.seller_level and "platinum" in seller.seller_level.lower():
        score += 2
    if seller.listing_count >= 3:
        score += 1
    if score >= 5:
        return "🔥 高"
    if score >= 3:
        return "⭐ 中"
    return "普通"


def suggest_action(seller: Seller) -> str:
    parts = []
    if seller.likely_chinese:
        parts.append("中文沟通, 加微信")
    else:
        parts.append("先发英文/西语开发信, 评估转化")
    if seller.is_official_store:
        parts.append("官方店, 大概率有自有品牌, 提 ODM/独家")
    if (seller.sold_total or 0) >= 1000:
        parts.append("销量大, 优先")
    if seller.registration_year and seller.registration_year >= 2023:
        parts.append("新卖家, 抢签代理窗口期")
    return "; ".join(parts)


def write_excel(
    listings: list[Listing],
    sellers: dict[str, Seller],
    output_dir: Path,
):
    output_dir.mkdir(parents=True, exist_ok=True)

    # ---------- sellers.xlsx ----------
    wb = Workbook()
    ws = wb.active
    ws.title = "Sellers"
    headers = [
        "优先级", "中国卖家?", "中国卖家信号", "卖家昵称",
        "累计销量", "信用等级", "官方店?",
        "国家", "州/省", "城市", "完整地址", "注册年份",
        "我们抓到的商品数", "样例商品标题", "店铺主页", "触达建议",
    ]
    ws.append(headers)
    style_header(ws)

    sorted_sellers = sorted(
        sellers.values(),
        key=lambda s: (
            not s.likely_chinese,         # 中国卖家在前
            -(s.sold_total or 0),
            -s.listing_count,
        ),
    )
    for sl in sorted_sellers:
        prio = priority_for(sl)
        row = [
            prio,
            "是" if sl.likely_chinese else "",
            ", ".join(sl.cn_signals),
            sl.nickname,
            sl.sold_total,
            sl.seller_level,
            "是" if sl.is_official_store else "",
            sl.country, sl.state, sl.city, sl.address_full,
            sl.registration_year,
            sl.listing_count,
            " | ".join(sl.sample_titles),
            sl.profile_url,
            suggest_action(sl),
        ]
        ws.append(row)
        # 着色
        last = ws.max_row
        if prio.startswith("🔥"):
            for c in ws[last]:
                c.fill = HOT_FILL
        elif prio.startswith("⭐"):
            for c in ws[last]:
                c.fill = WARM_FILL
    autosize(ws)
    ws.freeze_panes = "A2"

    seller_path = output_dir / "sellers.xlsx"
    wb.save(seller_path)
    print(f"[OUT] {seller_path}  ({len(sellers)} 卖家)")

    # ---------- listings.xlsx ----------
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "Listings"
    headers2 = [
        "搜索关键词", "商品ID", "标题", "价格", "货币", "销量",
        "卖家昵称", "卖家所在地", "商品链接",
    ]
    ws2.append(headers2)
    style_header(ws2)
    for l in listings:
        ws2.append([
            l.keyword, l.item_id, l.title, l.price, l.currency,
            l.sold, l.seller_nickname, l.location, l.url,
        ])
    autosize(ws2)
    ws2.freeze_panes = "A2"
    listing_path = output_dir / "listings.xlsx"
    wb2.save(listing_path)
    print(f"[OUT] {listing_path}  ({len(listings)} 商品)")


# ----------------------------------------------------------------------
# CLI
# ----------------------------------------------------------------------
def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="MercadoLibre 卖家爬虫 (爱骑仕跨境部)")
    p.add_argument("--config", default="config.yaml", help="配置文件路径")
    p.add_argument("--keyword", help="只抓单个关键词 (覆盖 config)")
    p.add_argument("--max-pages", type=int, help="每词最多翻多少页 (覆盖 config)")
    p.add_argument("--auto-proxy", action="store_true",
                   help="自动探测本机代理端口 (Clash/v2rayN等)")
    p.add_argument("--proxy", help="手动指定代理, 如 http://127.0.0.1:7890")
    return p.parse_args()


def main():
    args = parse_args()
    cfg_path = Path(args.config)
    if not cfg_path.exists():
        print(f"找不到 {cfg_path}", file=sys.stderr)
        sys.exit(1)
    with cfg_path.open(encoding="utf-8") as f:
        cfg = yaml.safe_load(f)

    print(f"[CFG] site={cfg['site']}, 关键词数={len(cfg.get('keywords', []))}")

    # 代理处理: --proxy > --auto-proxy > 已有环境变量 > 没代理(警告)
    proxy_url: Optional[str] = None
    if args.proxy:
        proxy_url = args.proxy
        print(f"[PROXY] 手动指定: {proxy_url}")
    elif args.auto_proxy:
        proxy_url = detect_proxy()
        if not proxy_url:
            print("[FATAL] 自动探测失败。检查你的 VPN/Clash/v2rayN 是否已开启,")
            print("       或手动用 --proxy http://127.0.0.1:端口号 指定")
            sys.exit(1)
    elif os.environ.get("HTTPS_PROXY") or os.environ.get("HTTP_PROXY"):
        proxy_url = os.environ.get("HTTPS_PROXY") or os.environ.get("HTTP_PROXY")
        print(f"[PROXY] 沿用环境变量: {proxy_url}")
    else:
        print("[PROXY] ⚠️ 未设置代理。如果你在中国大陆, MercadoLibre 一定会拦截。")
        print("       建议加 --auto-proxy 让脚本自动探测, 或 --proxy http://...")

    # 把代理塞进环境变量, 这样所有 requests 都自动走它
    if proxy_url:
        os.environ["HTTPS_PROXY"] = proxy_url
        os.environ["HTTP_PROXY"] = proxy_url
        # 验证代理 IP 真的能打开 MercadoLibre
        if not verify_proxy_works_for_ml(proxy_url, cfg["site"]):
            print("\n[FATAL] 代理可联网但 MercadoLibre 拦截了你的出口 IP。")
            print("        VPN 节点切到 智利/墨西哥/美国/巴西, 然后重跑。")
            sys.exit(1)

    listings, sellers = crawl(cfg, args)

    if not listings:
        print("\n没抓到任何商品。可能原因:")
        print("  1) IP 被 MercadoLibre 拦截 (开 VPN 或设置 HTTPS_PROXY)")
        print("  2) 关键词在该站没有结果 (改 config.yaml)")
        return

    output_dir = Path(cfg.get("output_dir", "output"))
    write_excel(listings, sellers, output_dir)
    print(f"\n[DONE] 完整结果在 {output_dir.resolve()}/")


if __name__ == "__main__":
    main()
